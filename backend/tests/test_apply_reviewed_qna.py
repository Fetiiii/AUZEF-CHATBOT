"""Reviewed-QnA workbook apply: colour audit, deterministic reconciliation, safe atomic apply."""
from __future__ import annotations

import pytest

from scripts import apply_reviewed_qna as ar

RED = ar.RED
GREEN = ar.GREEN


def _workbook(tmp_path, rows):
    """rows: (question, answer, tags, qna_colour, [(query, colour), ...])."""
    import openpyxl
    from openpyxl.styles import PatternFill

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append([None, "question", "answer", "tags"] + [f"query_{i}" for i in range(1, 21)])
    for r, (question, answer, tags, colour, queries) in enumerate(rows, start=2):
        ws.cell(r, 1, "reviewer")
        for col, value in ((2, question), (3, answer), (4, tags)):
            ws.cell(r, col, value).fill = PatternFill("solid", fgColor=colour)
        for i, (text, qcolour) in enumerate(queries):
            ws.cell(r, 5 + i, text).fill = PatternFill("solid", fgColor=qcolour)
    path = tmp_path / "reviewed.xlsx"
    wb.save(path)
    return str(path)


def _seed(db, rows):
    from core.database import QnA, QnAQuery

    ids = {}
    for key, question, answer, queries in rows:
        row = QnA(question_text=question, answer_text=answer, status=1)
        db.add(row)
        db.flush()
        for text in queries:
            db.add(QnAQuery(qna_id=row.id, query_text=text))
        ids[key] = row.id
    db.commit()
    return ids


# Reviewed workbook: an approved row whose answer and queries changed, a row whose
# question was edited (paired only via decisions), an exact red row and a red row
# whose answer differs (removal approved via decisions).
WORKBOOK = [
    ("Kayıt nasıl yapılır?", "Yeni cevap: AUZEF Öğrenme Yönetim Sistemi.", "", GREEN,
     [("kayıt", GREEN), ("başvuru iptali", RED), ("kayıt olmak istiyorum", GREEN)]),
    ("Neden 239,5 TL fazla ödüyorum?", "Katkı payı 239,5 TL.", "", GREEN, [("fazla ücret", GREEN)]),
    ("Eski soru?", "Eski cevap.", "", RED, [("eski", RED)]),
    ("Değişmiş kırmızı?", "Yeni metin.", "", RED, []),
]
DB_ROWS = [
    ("kept", "Kayıt nasıl yapılır?", "Eski cevap: AKSİS.", ["kayıt", "başvuru iptali", "kayıt olmak"]),
    ("edited", "Neden 86 TL fazla ödüyorum?", "Katkı payı 86 TL.", ["fazla ücret"]),
    ("red_exact", "Eski soru?", "Eski cevap.", ["eski"]),
    ("red_diff", "Değişmiş kırmızı?", "Eski metin.", []),
]


def _decisions(db, ids, workbook_path, *, sync=True, approve_red=True):
    snap = ar.db_snapshot(db)
    return {
        "workbook_sha256": ar.file_sha256(workbook_path),
        "sync_reviewed_content": sync,
        "approved_red_content_difference": [{"excel_row": 5, "db_id": ids["red_diff"]}] if approve_red else [],
        "question_edits": [{
            "excel_row": 3, "db_id": ids["edited"],
            "old_question_sha256": ar.raw_hash(snap[ids["edited"]]["question"]),
            "reviewed_question_sha256": ar.raw_hash("Neden 239,5 TL fazla ödüyorum?"),
        }],
    }


def _plan(db, rows, decisions):
    snap = ar.db_snapshot(db)
    rec = ar.reconcile(rows, snap, decisions)
    return rec, ar.build_plan(rows, rec, snap, decisions)


@pytest.fixture()
def seeded(db, tmp_path):
    ids = _seed(db, DB_ROWS)
    path = _workbook(tmp_path, WORKBOOK)
    rows, _ = ar.load_workbook_rows(path)
    return ids, path, rows


# ── audit ───────────────────────────────────────────────────────────────────


def test_workbook_colours_are_read_from_fill(tmp_path):
    rows, summary = ar.load_workbook_rows(_workbook(tmp_path, WORKBOOK))
    assert summary["qna_rows"] == 4 and summary["red_qna"] == 2
    assert summary["filled_queries"] == 5 and summary["red_queries_total"] == 2
    assert summary["red_queries_in_red_qna"] == 1 and summary["red_queries_in_kept_qna"] == 1
    assert summary["active_queries_after_cleanup"] == 3
    assert summary["anomalies"] == []


def test_reviewed_dataset_drops_red_qna_and_red_queries_and_duplicates(tmp_path):
    rows, _ = ar.load_workbook_rows(_workbook(tmp_path, [
        ("Soru?", "Cevap.", "", GREEN, [("a", GREEN), ("A ", GREEN), ("b", RED)]),
        ("Kırmızı?", "X.", "", RED, [("c", RED)]),
    ]))
    dataset, duplicates = ar.reviewed_dataset(rows)
    assert [d["question"] for d in dataset] == ["Soru?"]
    assert dataset[0]["queries"] == ["a"] and dataset[0]["review_status"] == "approved"
    assert len(duplicates) == 1


def test_normalization_is_deterministic_and_turkish_aware():
    assert ar.normalize("  İSTANBUL  Üniversitesi ") == ar.normalize("istanbul üniversitesi")
    assert ar.normalize("KAYIT") == "kayıt"
    assert ar.text_hash("a  b") == ar.text_hash("a b")


# ── without approvals the conservative behaviour is kept ────────────────────


def test_without_decisions_content_differences_and_edits_block(db, seeded):
    ids, path, rows = seeded
    rec, plan = _plan(db, rows, None)
    reasons = " ".join(b.get("reason", "") for b in plan["blockers"])
    assert "not approved for removal" in reasons          # red with content difference
    assert "kept QnA is UNMATCHED" in reasons            # edited question, no pairing
    assert "sync is not approved" in reasons             # changed answer


def test_decisions_for_a_different_workbook_are_rejected(db, seeded, tmp_path, monkeypatch):
    import json

    ids, path, rows = seeded
    decisions = _decisions(db, ids, path)
    decisions["workbook_sha256"] = "0" * 64
    target = tmp_path / "decisions.json"
    target.write_text(json.dumps(decisions), encoding="utf-8")
    with pytest.raises(SystemExit, match="different workbook"):
        ar.main(["--input", path, "--decisions", str(target), "--dry-run"])


# ── full apply ──────────────────────────────────────────────────────────────


def test_plan_covers_answer_question_queries_and_soft_delete(db, seeded):
    ids, path, rows = seeded
    rec, plan = _plan(db, rows, _decisions(db, ids, path))
    assert plan["blockers"] == []
    s = plan["summary"]
    assert (s["deactivate_qna"], s["update_answer"], s["update_question"]) == (2, 2, 1)
    assert (s["delete_queries"], s["insert_queries"]) == (2, 1)   # red + replaced; new text
    report = ar.plan_report(plan)
    update = next(o for o in report if o["op"] == "update_qna" and o["db_id"] == ids["edited"])
    assert update["old_question_sha256"] == ar.raw_hash("Neden 86 TL fazla ödüyorum?")
    assert update["reviewed_question_sha256"] == ar.raw_hash("Neden 239,5 TL fazla ödüyorum?")
    assert "old_answer_sha256" in update and "reviewed_answer_sha256" in update


def test_apply_syncs_workbook_and_second_run_is_a_no_op(db, seeded):
    from core.database import QnA, QnAQuery

    ids, path, rows = seeded
    decisions = _decisions(db, ids, path)
    rec, plan = _plan(db, rows, decisions)
    ar.apply_plan(db, plan)
    db.commit()
    db.expire_all()

    assert db.get(QnA, ids["red_exact"]).status == 0 and db.get(QnA, ids["red_diff"]).status == 0
    edited = db.get(QnA, ids["edited"])
    assert edited.question_text == "Neden 239,5 TL fazla ödüyorum?"
    assert edited.answer_text == "Katkı payı 239,5 TL." and edited.updated_by == ar.ACTOR
    assert db.get(QnA, ids["kept"]).answer_text == "Yeni cevap: AUZEF Öğrenme Yönetim Sistemi."
    kept_queries = sorted(q.query_text for q in db.query(QnAQuery).filter_by(qna_id=ids["kept"]))
    assert kept_queries == ["kayıt", "kayıt olmak istiyorum"]

    check = ar.verify(db, rows, rec)
    assert set(check["mismatch_counts"].values()) == {0}
    assert check["active_qna"] == 2 and check["active_queries"] == 3

    rec2, again = _plan(db, rows, decisions)
    assert again["blockers"] == [] and again["summary"]["total_mutations"] == 0


def test_unexpected_old_value_fails_before_writing(db, seeded):
    from core.database import QnA

    ids, path, rows = seeded
    rec, plan = _plan(db, rows, _decisions(db, ids, path))
    db.get(QnA, ids["kept"]).answer_text = "Başka biri düzenledi."    # concurrent admin edit
    db.commit()
    with pytest.raises(ar.ReviewedApplyError, match="expected old state"):
        ar.apply_plan(db, plan)
    db.rollback()
    db.expire_all()
    assert db.get(QnA, ids["red_exact"]).status == 1          # nothing was written


def test_error_mid_apply_rolls_back_everything(db, seeded, monkeypatch):
    from core.database import QnA, QnAQuery

    ids, path, rows = seeded
    rec, plan = _plan(db, rows, _decisions(db, ids, path))
    real = ar.apply_plan

    def failing(db_, plan_):
        real(db_, {**plan_, "summary": {**plan_["summary"], "delete_queries": -1}})

    with pytest.raises(ar.ReviewedApplyError, match="row counts differ"):
        failing(db, plan)
    db.rollback()
    db.expire_all()
    assert db.get(QnA, ids["red_exact"]).status == 1
    assert db.get(QnA, ids["kept"]).answer_text == "Eski cevap: AKSİS."
    assert db.query(QnAQuery).filter_by(qna_id=ids["kept"]).count() == 3


def test_decision_with_wrong_old_hash_blocks(db, seeded):
    ids, path, rows = seeded
    decisions = _decisions(db, ids, path)
    decisions["question_edits"][0]["old_question_sha256"] = "f" * 64
    rec, plan = _plan(db, rows, decisions)
    assert any("neither the approved old" in b.get("reason", "") for b in plan["blockers"])


def test_apply_requires_matching_plan_fingerprint_and_local_confirmation(db, seeded, tmp_path):
    import json

    ids, path, rows = seeded
    target = tmp_path / "decisions.json"
    target.write_text(json.dumps(_decisions(db, ids, path)), encoding="utf-8")
    base = ["--input", path, "--decisions", str(target), "--apply"]
    with pytest.raises(SystemExit, match="confirm-local-dev"):
        ar.main(base + ["--plan-fingerprint", "x"])
    with pytest.raises(SystemExit, match="plan-fingerprint does not match"):
        ar.main(base + ["--confirm-local-dev", "--plan-fingerprint", "x"])


def test_non_local_database_is_refused(monkeypatch):
    monkeypatch.setenv("ADMIN_DATABASE_URL", "postgresql://u:p@db-admin.prod.example:5432/auzef")
    with pytest.raises(SystemExit, match="not local/dev"):
        ar.assert_local_database()


def test_tag_difference_is_a_blocker_not_a_silent_change(db, seeded, tmp_path):
    ids, _, _ = seeded
    path = _workbook(tmp_path, [("Kayıt nasıl yapılır?", "Eski cevap: AKSİS.", "yeni-etiket", GREEN,
                                 [("kayıt", GREEN)])])
    rows, _ = ar.load_workbook_rows(path)
    rec, plan = _plan(db, rows, None)
    assert any("tag difference" in b.get("reason", "") for b in plan["blockers"])
